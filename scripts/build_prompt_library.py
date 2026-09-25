# -*- coding: utf-8 -*-
"""
Gera content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx — v2.

Reescrito a partir de prompts reais extraídos do concorrente (Studio
Pablita), aplicando os padrões observados: física de equilíbrio/gravidade,
vocabulário real de câmera e iluminação de estúdio, proteções explícitas
contra os erros clássicos de geração por IA (duplicação, flutuação,
reinterpretação de acabamento/escala, correntes/curvas simétricas demais),
composição sem rosto nas fotos humanizadas, sombra diagonal como assinatura
visual, e auto-checagem final antes de considerar a imagem pronta.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
ASPECTO_PADRAO = "3:4 (vertical)"

# ---------------------------------------------------------------------------
# Blocos fixos
# ---------------------------------------------------------------------------

FIDELIDADE = (
    "Utilize exclusivamente a joia da fotografia anexada como referência do "
    "produto. Preserve com fidelidade absoluta: formato, proporções, "
    "espessura, acabamento, brilho, textura, cor do metal, pedras (corte, "
    "tamanho e posição), cravação, fechos e elos, e todos os demais "
    "detalhes originais da peça. Não redesenhe, reinterprete, estilize, "
    "complete, corrija ou invente nenhuma característica da joia — "
    "reproduza exatamente o que existe na imagem original. Não utilize a "
    "foto anexada como referência de cenário, iluminação, enquadramento ou "
    "composição — ela serve apenas para identificar a peça e seus "
    "detalhes exatos; remova completamente quaisquer objetos, mãos ou "
    "fundos presentes na foto original."
)

FECHO_UNIVERSAL = (
    "Qualidade: altíssima resolução, nitidez máxima na joia, excelente "
    "alcance dinâmico, texturas naturais, reflexos fisicamente corretos, "
    "cores absolutamente fiéis, acabamento fotográfico profissional. "
    "Eliminar completamente aparência de renderização 3D, ruído digital, "
    "granulado, halos ou qualquer aspecto artificial típico de imagem "
    "gerada por IA. A peça não deve parecer flutuando, tombada ou fora de "
    "escala — preserve rigorosamente suas proporções e comportamento "
    "físico real. Antes de finalizar, confirme que a peça gerada "
    "corresponde exatamente, em quantidade e em detalhes, à peça da "
    "imagem original. Formato final da imagem: {aspecto}."
)

# ---------------------------------------------------------------------------
# Física/composição por tipo de peça — usada nas poses "still" (sem modelo)
# ---------------------------------------------------------------------------

TIPO_COMPOSICAO_STILL = {
    "Brinco": (
        "Deixe um brinco apoiado em pé, em ângulo de aproximadamente 45° "
        "em relação à câmera, e o outro brinco deitado naturalmente sobre "
        "a superfície ao lado. Caso a peça tenha tarraxa, mantenha uma "
        "tarraxa presa em cada brinco — nunca deixe a tarraxa solta ou "
        "separada da peça."
    ),
    "Colar": (
        "Apoie o colar completamente sobre a superfície, deixando a "
        "corrente cair em curvas amplas e orgânicas, como o peso natural "
        "do metal realmente se comportaria — nunca reproduza o formato "
        "exato da foto original. Nunca crie círculos perfeitos, ovais, "
        "corações, gotas ou curvas espelhadas/simétricas dos dois lados; "
        "cada lado da corrente deve ter um desenho diferente, com raios e "
        "comprimentos de curva variados. Se houver pingente, ele deve ser "
        "o ponto focal da composição."
    ),
    "Choker/Gargantilha": (
        "Apoie a peça em uma curva aberta e levemente assimétrica sobre a "
        "superfície, sugerindo o formato do pescoço sem formar um círculo "
        "perfeito ou fechado."
    ),
    "Colar Longo": (
        "Apoie a corrente formando curvas amplas, orgânicas e "
        "assimétricas — nunca curvas espelhadas ou geométricas. Deixe o "
        "pingente (ou uma das extremidades) como ponto focal da "
        "composição, com o restante se estendendo de forma natural, como "
        "se tivesse acabado de cair sobre a superfície."
    ),
    "Pingente": (
        "Posicione o pingente como protagonista absoluto, com a corrente "
        "caindo ao redor em curvas orgânicas e assimétricas — nunca "
        "círculos, ovais ou formas espelhadas dos dois lados."
    ),
    "Corrente Masculina": (
        "Apoie a corrente em uma curva reta a levemente ondulada, "
        "transmitindo peso e robustez — nunca curvas perfeitamente "
        "simétricas ou decorativas demais."
    ),
    "Pulseira": (
        "Apoie a pulseira formando uma curva aberta e natural sobre a "
        "superfície, como se tivesse caído naturalmente — nunca um "
        "círculo perfeitamente fechado."
    ),
    "Bracelete/Bangle": (
        "Posicione o bracelete em pé, com o eixo central perpendicular à "
        "superfície (aproximadamente 90°) e inclinação máxima de 3° — o "
        "suficiente apenas para evitar aparência artificial de "
        "renderização. O centro de gravidade deve estar exatamente sobre "
        "o ponto de apoio, sem parecer tombado ou flutuando. A peça deve "
        "ocupar cerca de 65% da altura da imagem."
    ),
    "Anel": (
        "Posicione o anel em pé, levemente inclinado, mostrando "
        "claramente o aro e o detalhe frontal (pedra ou desenho "
        "principal). A peça deve parecer fisicamente apoiada e "
        "equilibrada, nunca flutuando ou tombada."
    ),
    "Aliança": (
        "Posicione as duas alianças em pé, lado a lado, uma levemente "
        "sobreposta à outra, ambas fisicamente equilibradas sobre a "
        "superfície — nunca flutuando."
    ),
    "Tornozeleira": (
        "Apoie a peça formando uma curva aberta e natural, como se "
        "tivesse caído sobre a superfície — nunca um círculo "
        "perfeitamente fechado."
    ),
    "Broche": (
        "Apoie o broche sobre uma prega de tecido, mostrando claramente o "
        "pino e o relevo frontal da peça, sem parecer flutuando."
    ),
    "Conjunto (colar + brinco)": (
        "Utilize exatamente as peças presentes na foto anexada, na mesma "
        "quantidade — nunca adicione, remova ou duplique nenhuma joia. Se "
        "não houver um tipo de peça na foto, não crie esse tipo de peça. "
        "Organize cada peça de acordo com seu próprio comportamento "
        "físico (colar em curvas orgânicas e assimétricas; brincos "
        "próximos entre si, mostrando a face principal), distribuindo os "
        "elementos em diferentes regiões da composição para um equilíbrio "
        "visual sofisticado, porém não perfeitamente simétrico."
    ),
    "Relógio": (
        "Posicione o relógio em pé, mostrando o mostrador de frente, com "
        "a pulseira formando uma curva aberta e natural ao redor — "
        "fisicamente equilibrado, nunca flutuando ou tombado."
    ),
}

# Parte do corpo mostrada nas poses humanizadas/lifestyle (sempre sem rosto)
TIPO_PARTE_CORPO = {
    "Brinco": "a orelha e a lateral do rosto (apenas esses elementos — não mostrar o restante do rosto)",
    "Colar": "o colo, as clavículas e a parte inferior do pescoço, terminando imediatamente abaixo da mandíbula",
    "Choker/Gargantilha": "o pescoço",
    "Colar Longo": "o colo e o decote",
    "Pingente": "o colo, próximo ao centro do peito",
    "Corrente Masculina": "o peito e a base do pescoço, com estilo masculino",
    "Pulseira": "o pulso",
    "Bracelete/Bangle": "o pulso e o antebraço",
    "Anel": "a mão e os dedos",
    "Aliança": "as mãos entrelaçadas, próximas ao dedo anelar",
    "Tornozeleira": "o tornozelo e o pé",
    "Broche": "a lapela de um blazer ou prega de tecido",
    "Conjunto (colar + brinco)": "o pescoço, o colo e a orelha",
    "Relógio": "o pulso",
}

GENERO_TIPO = {
    "Brinco": "feminino", "Colar": "feminino", "Choker/Gargantilha": "feminino",
    "Colar Longo": "feminino", "Pingente": "feminino", "Corrente Masculina": "masculino",
    "Pulseira": "feminino", "Bracelete/Bangle": "feminino", "Anel": "feminino",
    "Aliança": "unissex", "Tornozeleira": "feminino", "Broche": "feminino",
    "Conjunto (colar + brinco)": "feminino", "Relógio": "unissex",
}

PERFIS_SUGERIDOS = {
    "feminino": "Mulher jovem (18-25), pele clara | Mulher adulta (30-45), pele morena | Mulher adulta (30-45), pele negra | Mulher idosa (60+), pele clara",
    "masculino": "Homem jovem (20-30), pele morena | Homem adulto (35-50), pele negra | Homem adulto (35-50), pele clara",
    "unissex": "Mulher adulta (30-45), pele morena | Homem adulto (35-50), pele clara | Mulher idosa (60+), pele negra",
}

TEXTURAS_SUPERFICIE = ["mármore branco", "madeira clara", "pedra calcária clara", "linho bege"]

# ---------------------------------------------------------------------------
# Especificação técnica por pose
# ---------------------------------------------------------------------------

def tech_still_branco():
    return (
        "Fotografia still de produto (packshot) profissional para "
        "e-commerce. Fundo branco puro (RGB 255,255,255), infinito, sem "
        "linha de horizonte, sem gradientes, sem textura e sem qualquer "
        "elemento decorativo. Iluminação de estúdio: uma fonte de luz "
        "principal ampla e difusa, uma luz de preenchimento suave do lado "
        "oposto para controlar sombras, e uma luz de contorno sutil para "
        "destacar a silhueta. Reflexos no metal limpos, contínuos e "
        "fisicamente corretos, sem brilhos estourados. Sombra de contato "
        "única, suave e difusa, com leve gradiente natural logo abaixo da "
        "peça. Lente equivalente entre 85mm e 105mm, foco perfeitamente "
        "nítido em toda a peça, sem distorção de perspectiva."
    )

def tech_still_sombra_editorial():
    return (
        "Fotografia editorial de joias com estética minimalista, "
        "sofisticada e contemporânea, semelhante a campanhas de marcas de "
        "luxo. A peça deve estar apoiada sobre um tecido branco ou "
        "off-white premium, fosco, encorpado, com textura delicada "
        "semelhante a crepe de seda estruturado, moldado à mão em ondas "
        "amplas e curvas orgânicas — nunca padrões repetitivos, simetria "
        "perfeita ou dobras geométricas/espirais matematicamente exatas "
        "(isso denuncia aparência de imagem gerada por IA). Câmera em "
        "ângulo de aproximadamente 45°, lente macro equivalente a "
        "85-100mm, enquadramento em close-up com a peça ocupando cerca da "
        "metade da imagem. Iluminação natural extremamente suave e "
        "difusa, como luz de uma grande janela lateral, criando sombras "
        "delicadas entre as dobras do tecido. Paleta restrita a branco "
        "quente, marfim e creme suave. Não adicionar flores, folhas, "
        "pedras decorativas, madeira, embalagens, mãos, pessoas ou "
        "qualquer elemento decorativo."
    )

def tech_still_textura(idx):
    textura = TEXTURAS_SUPERFICIE[idx % len(TEXTURAS_SUPERFICIE)]
    return (
        f"Fotografia still editorial sobre uma superfície de {textura}, "
        "de formato orgânico e bordas irregulares, apoiada sobre um "
        "tecido acetinado em tom bege claro e quente. Câmera em "
        "perspectiva oblíqua, entre 45° e 55° acima da superfície — nunca "
        "totalmente frontal nem top-down perfeitamente perpendicular; a "
        "composição pode aparecer levemente rotacionada dentro do quadro. "
        "Iluminação profissional de estúdio inspirada em luz solar suave "
        "entrando lateralmente, criando sombras pequenas e realistas que "
        "mostram a peça fisicamente apoiada sobre a superfície. Contraste "
        "moderado, cena clara, quente e sofisticada — evitar sombras "
        "muito escuras ou iluminação dramática."
    )

def tech_still_humanizado():
    return (
        "Fotografia editorial extremamente realista para catálogo premium "
        "de joias, produzida como uma campanha de joalheria de luxo. "
        "Enquadramento fechado, mostrando apenas [PARTE DO CORPO] — nunca "
        "mostrar rosto, olhos, boca, nariz, cabelo ou outras partes do "
        "corpo além das indicadas. A pele deve ter aparência extremamente "
        "natural, saudável e uniforme, com tom: [PERFIL]. Criar uma única "
        "faixa de sombra diagonal, contínua e bem definida, inclinada "
        "aproximadamente 45°, atravessando a composição como luz entrando "
        "por uma janela — a sombra nunca deve cruzar ou escurecer a peça, "
        "que deve permanecer completamente iluminada e em destaque. "
        "Vestuário minimalista em tecido acetinado ou seda fosca, em tom "
        "off-white, extremamente discreto para não competir com a joia. "
        "Lente equivalente entre 85mm e 105mm, profundidade de campo "
        "rasa — a pele pode ficar levemente desfocada, mas a peça deve "
        "permanecer perfeitamente nítida. Fundo neutro (cinza quente "
        "muito suave ou bege acinzentado), sem textura, objetos ou "
        "elementos decorativos. Não adicionar outras joias, acessórios ou "
        "maquiagem chamativa."
    )

def tech_lifestyle():
    return (
        "Fotografia lifestyle editorial para redes sociais e e-commerce. "
        "Enquadramento fechado o suficiente para não mostrar o rosto da "
        "modelo — foco em [PARTE DO CORPO] usando a peça de forma "
        "natural, em um gesto espontâneo do dia a dia, compatível com o "
        "tipo de peça. Iluminação natural suave, como luz de uma janela "
        "lateral ampla, sem fonte artificial. Fundo desfocado "
        "(profundidade de campo rasa), sugerindo um ambiente interno "
        "elegante e minimalista, sem elementos que disputem atenção com a "
        "joia. Tons neutros e quentes (branco, off-white, bege). A peça "
        "deve permanecer perfeitamente nítida e iluminada, com reflexos "
        "naturais e fisicamente corretos no metal."
    )

POSES_ORDEM = [
    "Still Branco (Packshot)",
    "Still Humanizado",
    "Still com Sombra Editorial",
    "Still em Superfície Texturizada",
    "Lifestyle com Modelo",
]

TEM_MODELO = {
    "Still Branco (Packshot)": "Não",
    "Still Humanizado": "Sim (parte do corpo, sem rosto)",
    "Still com Sombra Editorial": "Não",
    "Still em Superfície Texturizada": "Não",
    "Lifestyle com Modelo": "Sim (parte do corpo, sem rosto)",
}

# ---------------------------------------------------------------------------
# Datas comemorativas
# ---------------------------------------------------------------------------

DATAS_TEMA = {
    "Dia das Mães": {
        "paleta": "tons de rosa suave e nude",
        "props": "flores (rosas) desfocadas ao fundo, tecido de seda rosa-claro",
        "clima": "afetivo, delicado e caloroso",
    },
    "Dia dos Namorados": {
        "paleta": "tons de vermelho e vinho",
        "props": "pétalas de rosa vermelha, velas acesas desfocadas ao fundo",
        "clima": "romântico, intimista, iluminação quente e noturna",
    },
    "Dia das Crianças": {
        "paleta": "tons pastéis vibrantes (amarelo, azul-bebê, rosa-claro)",
        "props": "pequenos balões desfocados ao fundo, ambiente lúdico",
        "clima": "leve, alegre e descontraído",
    },
    "Dia dos Pais": {
        "paleta": "tons de azul-marinho e cinza-grafite",
        "props": "madeira escura, textura de couro ao fundo desfocado",
        "clima": "sóbrio, elegante e masculino",
    },
    "Dia Internacional da Mulher": {
        "paleta": "tons de roxo e violeta",
        "props": "flor de mimosa (amarela) desfocada ao fundo",
        "clima": "moderno, empoderado e sofisticado",
    },
    "Black Friday": {
        "paleta": "preto e dourado",
        "props": "fundo preto fosco, iluminação dramática lateral, reflexo dourado sutil",
        "clima": "luxuoso, de destaque, alto contraste",
    },
    "Natal": {
        "paleta": "tons de vermelho, dourado e verde",
        "props": "ramos verdes, pinha desfocada, luzes de natal desfocadas ao fundo (efeito bokeh dourado)",
        "clima": "aconchegante, festivo e brilhante",
    },
    "Consciência Negra": {
        "paleta": "tons terrosos e dourados",
        "props": "tecido com estampa de padronagem africana desfocado ao fundo",
        "clima": "de valorização, orgulho e sofisticação",
    },
}

TIPOS_DATA = ["Brinco", "Colar", "Anel", "Pulseira"]
POSES_DATA = ["Still com Tema", "Lifestyle com Tema"]

# ---------------------------------------------------------------------------
# Montagem dos prompts
# ---------------------------------------------------------------------------

def slug(txt):
    txt = txt.lower()
    repl = {
        "á": "a", "ã": "a", "â": "a", "à": "a", "é": "e", "ê": "e",
        "í": "i", "ó": "o", "ô": "o", "õ": "o", "ú": "u", "ç": "c",
        "(": "", ")": "", "/": "-", " ": "-",
    }
    for a, b in repl.items():
        txt = txt.replace(a, b)
    return txt

def montar_still(tipo, pose, textura_idx):
    cena = f"Composição: {TIPO_COMPOSICAO_STILL[tipo]}"
    if pose == "Still Branco (Packshot)":
        tech = tech_still_branco()
    elif pose == "Still com Sombra Editorial":
        tech = tech_still_sombra_editorial()
    else:  # Still em Superfície Texturizada
        tech = tech_still_textura(textura_idx)
    fecho = FECHO_UNIVERSAL.format(aspecto=ASPECTO_PADRAO)
    return "\n\n".join([FIDELIDADE, cena, tech, fecho])

def montar_humanizado_ou_lifestyle(tipo, pose):
    parte = TIPO_PARTE_CORPO[tipo]
    cena = f"Composição: mostrar apenas {parte}, evidenciando a peça de forma natural."
    tech = tech_still_humanizado() if pose == "Still Humanizado" else tech_lifestyle()
    tech = tech.replace("[PARTE DO CORPO]", parte)
    fecho = FECHO_UNIVERSAL.format(aspecto=ASPECTO_PADRAO)
    return "\n\n".join([FIDELIDADE, cena, tech, fecho])

def montar_prompt_evergreen(tipo, pose, textura_idx):
    if pose in ("Still Branco (Packshot)", "Still com Sombra Editorial", "Still em Superfície Texturizada"):
        return montar_still(tipo, pose, textura_idx)
    return montar_humanizado_ou_lifestyle(tipo, pose)

def montar_prompt_data(tipo, pose, data, tema):
    tema_txt = (
        f"Tema visual: {data}. Paleta de cores: {tema['paleta']}. "
        f"Elementos/props de cena: {tema['props']}. Clima geral: {tema['clima']}."
    )
    if pose == "Still com Tema":
        cena = f"Composição: {TIPO_COMPOSICAO_STILL[tipo]}"
        tech = tech_still_sombra_editorial()
    else:  # Lifestyle com Tema
        parte = TIPO_PARTE_CORPO[tipo]
        cena = f"Composição: mostrar apenas {parte}, evidenciando a peça de forma natural."
        tech = tech_still_humanizado().replace("[PARTE DO CORPO]", parte)
    fecho = FECHO_UNIVERSAL.format(aspecto=ASPECTO_PADRAO)
    return "\n\n".join([FIDELIDADE, cena, tema_txt, tech, fecho])

rows_evergreen = []
idx = 1
for tipo in TIPO_COMPOSICAO_STILL:
    genero = GENERO_TIPO[tipo]
    for ti, pose in enumerate(POSES_ORDEM):
        prompt = montar_prompt_evergreen(tipo, pose, ti)
        tem_modelo = TEM_MODELO[pose]
        perfis = PERFIS_SUGERIDOS[genero] if "Sim" in tem_modelo else ""
        nome_ref = f"{slug(tipo)}_{slug(pose)}"
        rows_evergreen.append([
            idx, tipo, pose, tem_modelo, perfis, nome_ref, prompt, ASPECTO_PADRAO, "A gerar", "", "",
        ])
        idx += 1

rows_datas = []
idx = 1
for data, tema in DATAS_TEMA.items():
    for tipo in TIPOS_DATA:
        genero = GENERO_TIPO[tipo]
        for pose in POSES_DATA:
            prompt = montar_prompt_data(tipo, pose, data, tema)
            tem_modelo = "Não" if pose == "Still com Tema" else "Sim (parte do corpo, sem rosto)"
            perfis = PERFIS_SUGERIDOS[genero] if tem_modelo.startswith("Sim") else ""
            nome_ref = f"{slug(data)}_{slug(tipo)}_{slug(pose)}"
            rows_datas.append([
                idx, data, tipo, pose, tem_modelo, perfis, nome_ref, prompt, ASPECTO_PADRAO, "A gerar", "", "",
            ])
            idx += 1

# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

wb = Workbook()

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name=FONT_NAME, size=10)
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="1F2937")
SUB_FONT = Font(name=FONT_NAME, size=10, italic=True, color="555555")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

def autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

ws0 = wb.active
ws0.title = "Instruções"
ws0.sheet_view.showGridLines = False
ws0["B2"] = "Photo Studio TPV — Biblioteca de Prompts (v2)"
ws0["B2"].font = TITLE_FONT
ws0["B3"] = "Reescrita com base em prompts reais do concorrente (Studio Pablita)"
ws0["B3"].font = SUB_FONT

instrucoes = [
    "",
    "O QUE MUDOU NA V2",
    "Os prompts foram reescritos incorporando técnicas reais observadas em prompts do "
    "concorrente: física de equilíbrio/gravidade das peças, vocabulário de câmera e "
    "iluminação de estúdio, proteções contra erros clássicos de geração por IA "
    "(duplicação, flutuação, correntes/curvas simétricas demais, reinterpretação de "
    "acabamento ou escala), composição sem rosto nas fotos humanizadas, sombra "
    "diagonal como assinatura visual, e auto-checagem final antes de considerar a "
    "imagem pronta.",
    "",
    "COMO USAR ESTA PLANILHA",
    "1. Abra a aba 'Evergreen' ou 'Datas Comemorativas'.",
    "2. Copie o texto da coluna 'Prompt Mestre' de uma linha.",
    "3. Se o prompt tiver [PERFIL], substitua por uma das opções de 'Perfis Sugeridos'.",
    "4. Cole no Gemini (ou ChatGPT), anexando a foto de uma peça genérica do catálogo.",
    "5. Gere a imagem. Se ficou boa, salve com o nome de 'Nome de Referência'.",
    "6. Suba no Supabase e marque a linha como 'aprovado' (ver README do projeto).",
    "",
    "LEGENDA",
    "Tem Modelo? — indica se a pose usa parte do corpo (sempre sem rosto).",
    "Perfis Sugeridos — variações de gênero/idade/tom de pele para gerar mais de uma "
    "vez o mesmo prompt, ampliando a diversidade de referências no catálogo.",
]

r = 5
for line in instrucoes:
    cell = ws0.cell(row=r, column=2, value=line)
    if line.isupper() and line != "":
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="1F2937")
    else:
        cell.font = CELL_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.row_dimensions[r].height = 28 if len(line) > 70 else 16
    r += 1

autofit(ws0, [3, 110])
for rr in range(5, r):
    ws0.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=10)

ws1 = wb.create_sheet("Evergreen")
headers = ["ID", "Tipo de Peça", "Pose/Enquadramento", "Tem Modelo?", "Perfis Sugeridos",
           "Nome de Referência", "Prompt Mestre", "Proporção", "Status",
           "Link da Imagem Gerada", "Observações"]
ws1.append(headers)
style_header(ws1, len(headers))
for row in rows_evergreen:
    ws1.append(row)

widths1 = [5, 20, 24, 22, 40, 34, 80, 16, 12, 26, 22]
autofit(ws1, widths1)
for r_i in range(2, ws1.max_row + 1):
    for c_i in range(1, len(headers) + 1):
        cell = ws1.cell(row=r_i, column=c_i)
        cell.font = CELL_FONT
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws1.row_dimensions[r_i].height = 130
ws1.freeze_panes = "A2"
ws1.auto_filter.ref = ws1.dimensions

ws2 = wb.create_sheet("Datas Comemorativas")
headers2 = ["ID", "Data Comemorativa", "Tipo de Peça", "Pose/Enquadramento", "Tem Modelo?",
            "Perfis Sugeridos", "Nome de Referência", "Prompt Mestre", "Proporção",
            "Status", "Link da Imagem Gerada", "Observações"]
ws2.append(headers2)
style_header(ws2, len(headers2))
for row in rows_datas:
    ws2.append(row)

widths2 = [5, 24, 18, 20, 22, 40, 38, 80, 16, 12, 26, 22]
autofit(ws2, widths2)
for r_i in range(2, ws2.max_row + 1):
    for c_i in range(1, len(headers2) + 1):
        cell = ws2.cell(row=r_i, column=c_i)
        cell.font = CELL_FONT
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.row_dimensions[r_i].height = 140
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = ws2.dimensions

out_path = "content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx"
wb.save(out_path)
print("OK:", out_path)
print("Evergreen rows:", len(rows_evergreen))
print("Datas rows:", len(rows_datas))
