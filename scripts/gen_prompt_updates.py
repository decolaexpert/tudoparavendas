# -*- coding: utf-8 -*-
"""
Gera supabase/prompt_updates/update_part_NN.sql a partir da planilha de
prompts — UPDATE que mexe SOMENTE em prompt_mestre, casado por
nome_referencia. Nunca toca em status/thumbnail_url, então fotos já
aprovadas no catálogo continuam intactas.

Uso (depois de rodar scripts/build_prompt_library.py):
    python3 scripts/gen_prompt_updates.py
"""
import math
import openpyxl

XLSX_PATH = "content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx"
OUT_DIR = "supabase/prompt_updates"
ROWS_PER_FILE = 8


def esc(value):
    return "'" + str(value).replace("'", "''") + "'"


def rows_from(ws, ref_col, prompt_col):
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append((row[idx[ref_col]], row[idx[prompt_col]]))
    return out


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    rows = rows_from(wb["Evergreen"], "Nome de Referência", "Prompt Mestre")
    rows += rows_from(wb["Datas Comemorativas"], "Nome de Referência", "Prompt Mestre")

    total_parts = math.ceil(len(rows) / ROWS_PER_FILE)
    for part in range(total_parts):
        chunk = rows[part * ROWS_PER_FILE : (part + 1) * ROWS_PER_FILE]
        lines = [
            f"-- Parte {part + 1}/{total_parts} -- atualiza SOMENTE prompt_mestre",
            "-- Nao mexe em status nem thumbnail_url -- fotos ja aprovadas ficam intactas",
            "",
        ]
        for nome_ref, prompt in chunk:
            lines.append(
                f"update public.reference_photos set prompt_mestre = {esc(prompt)} "
                f"where nome_referencia = {esc(nome_ref)};"
            )
        out_path = f"{OUT_DIR}/update_part_{part + 1:02d}.sql"
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")

    print(f"OK: {len(rows)} linhas em {total_parts} arquivos ({OUT_DIR}/update_part_NN.sql)")


if __name__ == "__main__":
    main()
