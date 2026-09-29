# -*- coding: utf-8 -*-
"""
Migração pontual v5: remove Relógio e a estrutura antiga de Datas
Comemorativas do banco, e insere a nova estrutura (Datas Comemorativas
reduzida a 5 fotos/data + a nova categoria Expositores).

Gera supabase/migration_v5/00_delete.sql (deletes) e
supabase/migration_v5/01..NN_insert.sql (inserts em chunks pequenos, pra
não esbarrar no limite de colagem do SQL Editor do Supabase).

Uso (depois de rodar scripts/build_prompt_library.py):
    python3 scripts/gen_migration_v5.py
"""
import math
import openpyxl

XLSX_PATH = "content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx"
OUT_DIR = "supabase/migration_v5"
ROWS_PER_FILE = 8


def esc(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def rows_from_datas(ws):
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append({
            "categoria": "data_comemorativa",
            "tipo_peca": row[idx["Tipo de Peça"]],
            "pose": row[idx["Pose/Enquadramento"]],
            "data_comemorativa": row[idx["Data Comemorativa"]],
            "nome_referencia": row[idx["Nome de Referência"]],
            "prompt_mestre": row[idx["Prompt Mestre"]],
            "perfis_sugeridos": row[idx["Perfis Sugeridos"]] or None,
            "aspecto": row[idx["Proporção"]],
            "ordem": row[idx["ID"]],
        })
    return out


def rows_from_expositores(ws):
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append({
            "categoria": "evergreen",
            "tipo_peca": row[idx["Tipo de Peça"]],
            "pose": row[idx["Pose/Enquadramento"]],
            "data_comemorativa": None,
            "nome_referencia": row[idx["Nome de Referência"]],
            "prompt_mestre": row[idx["Prompt Mestre"]],
            "perfis_sugeridos": row[idx["Perfis Sugeridos"]] or None,
            "aspecto": row[idx["Proporção"]],
            "ordem": row[idx["ID"]],
        })
    return out


def insert_stmt(r):
    cols = [
        "categoria", "tipo_peca", "pose", "data_comemorativa",
        "nome_referencia", "prompt_mestre", "perfis_sugeridos",
        "aspecto", "ordem", "status",
    ]
    values = [
        esc(r["categoria"]), esc(r["tipo_peca"]), esc(r["pose"]),
        esc(r["data_comemorativa"]), esc(r["nome_referencia"]),
        esc(r["prompt_mestre"]), esc(r["perfis_sugeridos"]),
        esc(r["aspecto"]), str(r["ordem"]), "'a_gerar'",
    ]
    return (
        f"insert into public.reference_photos ({', '.join(cols)}) "
        f"values ({', '.join(values)}) "
        f"on conflict (nome_referencia) do nothing;"
    )


def main():
    import os
    os.makedirs(OUT_DIR, exist_ok=True)

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    rows = rows_from_datas(wb["Datas Comemorativas"]) + rows_from_expositores(wb["Expositores"])

    delete_sql = [
        "-- Migração v5: remove Relógio (categoria descontinuada) e a",
        "-- estrutura antiga de Datas Comemorativas (nenhuma dessas linhas",
        "-- chegou a ser aprovada, então é seguro apagar e recriar).",
        "",
        "delete from public.reference_photos where nome_referencia like 'relogio\\_%' escape '\\';",
        "delete from public.reference_photos where categoria = 'data_comemorativa';",
        "",
    ]
    with open(f"{OUT_DIR}/00_delete.sql", "w", encoding="utf-8") as f:
        f.write("\n".join(delete_sql) + "\n")

    total_parts = math.ceil(len(rows) / ROWS_PER_FILE)
    for part in range(total_parts):
        chunk = rows[part * ROWS_PER_FILE : (part + 1) * ROWS_PER_FILE]
        lines = [
            f"-- Migração v5 -- insert parte {part + 1}/{total_parts}",
            "-- Datas Comemorativas (nova estrutura, 5 fotos/data) + Expositores",
            "",
        ]
        lines += [insert_stmt(r) for r in chunk]
        out_path = f"{OUT_DIR}/{part + 1:02d}_insert.sql"
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")

    print(f"OK: 00_delete.sql + {total_parts} arquivos de insert ({len(rows)} linhas) em {OUT_DIR}/")


if __name__ == "__main__":
    main()
