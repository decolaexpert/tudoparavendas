# -*- coding: utf-8 -*-
"""
Migração pontual v6: insere as duas novas categorias — Lifestyle (10
fotos, modelo genérica de rosto completo) e Peças em Você (10 fotos que
usam 2 fotos anexadas: a da assinante + a da joia). Não deleta nada —
são linhas totalmente novas, com nome_referencia inédito.

Gera supabase/migration_v6/01..NN_insert.sql em chunks pequenos, pra não
esbarrar no limite de colagem do SQL Editor do Supabase.

Uso (depois de rodar scripts/build_prompt_library.py):
    python3 scripts/gen_migration_v6.py
"""
import math
import os
import openpyxl

XLSX_PATH = "content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx"
OUT_DIR = "supabase/migration_v6"
ROWS_PER_FILE = 8


def esc(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def rows_from(ws):
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append({
            "categoria": "evergreen",
            "tipo_peca": row[idx["Tipo de Peça"]],
            "pose": row[idx["Pose/Enquadramento"]],
            "data_comemorativa": None,
            "nome_referencia": row[idx["Nome de Referência"]],
            "prompt_mestre": row[idx["Prompt Mestre"]],
            "perfis_sugeridos": row[idx["Perfis Sugeridos"]] or None,
            "aspecto": row[idx["Proporção"]],
            "ordem": row[idx["ID"]],
        })
    return out


def insert_stmt(r):
    cols = [
        "categoria", "tipo_peca", "pose", "data_comemorativa",
        "nome_referencia", "prompt_mestre", "perfis_sugeridos",
        "aspecto", "ordem", "status",
    ]
    values = [
        esc(r["categoria"]), esc(r["tipo_peca"]), esc(r["pose"]),
        esc(r["data_comemorativa"]), esc(r["nome_referencia"]),
        esc(r["prompt_mestre"]), esc(r["perfis_sugeridos"]),
        esc(r["aspecto"]), str(r["ordem"]), "'a_gerar'",
    ]
    return (
        f"insert into public.reference_photos ({', '.join(cols)}) "
        f"values ({', '.join(values)}) "
        f"on conflict (nome_referencia) do nothing;"
    )


def main():
    os.makedirs(OUT_DIR, exist_ok=True)

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    rows = rows_from(wb["Lifestyle"]) + rows_from(wb["Peças em Você"])

    total_parts = math.ceil(len(rows) / ROWS_PER_FILE)
    for part in range(total_parts):
        chunk = rows[part * ROWS_PER_FILE : (part + 1) * ROWS_PER_FILE]
        lines = [
            f"-- Migração v6 -- insert parte {part + 1}/{total_parts}",
            "-- Novas categorias: Lifestyle + Peças em Você",
            "",
        ]
        lines += [insert_stmt(r) for r in chunk]
        out_path = f"{OUT_DIR}/{part + 1:02d}_insert.sql"
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")

    print(f"OK: {total_parts} arquivos de insert ({len(rows)} linhas) em {OUT_DIR}/")


if __name__ == "__main__":
    main()
