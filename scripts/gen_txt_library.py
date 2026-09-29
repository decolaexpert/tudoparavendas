# -*- coding: utf-8 -*-
"""
Gera um .txt legível (pra copiar e colar no Gemini/ChatGPT) a partir de
content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx.

Uso (depois de rodar scripts/build_prompt_library.py):
    python3 scripts/gen_txt_library.py [caminho_de_saida.txt]
"""
import sys
import openpyxl

XLSX_PATH = "content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx"
VERSAO = "v5"
CHANGELOG = (
    "v5: remove a categoria Relógio; Datas Comemorativas passou a ter "
    "apenas 5 fotos por data (variando peça e cenário) em vez de 8; "
    "adiciona a nova aba Expositores (still de joia sobre suporte de "
    "exibição: busto, mão, orelha, correntes). Mantém a fidelidade de "
    "quantidade e a checagem de anatomia da v4."
)
SEP = "-" * 72
DSEP = "=" * 72
HSEP = "#" * 70


def rows_from(ws, extra_cols):
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {
            "tipo_peca": row[idx["Tipo de Peça"]],
            "pose": row[idx["Pose/Enquadramento"]],
            "nome_referencia": row[idx["Nome de Referência"]],
            "prompt_mestre": row[idx["Prompt Mestre"]],
            "perfis_sugeridos": row[idx["Perfis Sugeridos"]] or None,
            "aspecto": row[idx["Proporção"]],
        }
        for col in extra_cols:
            item[col] = row[idx[col]]
        out.append(item)
    return out


def render_item(lines, n, titulo, item):
    lines.append(SEP)
    lines.append(f"[{n}] {titulo}")
    lines.append(f"nome_referencia: {item['nome_referencia']}")
    lines.append(f"proporção sugerida: {item['aspecto']}")
    if item["perfis_sugeridos"]:
        lines.append(
            "perfis sugeridos (se tiver [PERFIL] no prompt, troque por um "
            f"destes): {item['perfis_sugeridos']}"
        )
    lines.append("")
    lines.append("PROMPT:")
    lines.append(item["prompt_mestre"])
    lines.append("")


def main():
    out_path = sys.argv[1] if len(sys.argv) > 1 else "Prompts_Gemini_Photo_Studio_TPV.txt"

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    evergreen = rows_from(wb["Evergreen"], [])
    datas = rows_from(wb["Datas Comemorativas"], ["Data Comemorativa"])
    expositores = rows_from(wb["Expositores"], [])

    lines = [
        f"PHOTO STUDIO TPV — PROMPTS PARA GERAR AS FOTOS-MESTRE NO GEMINI ({VERSAO})",
        DSEP,
        "",
        CHANGELOG,
        "",
        "Como usar:",
        "1. Copie o bloco 'PROMPT' de uma referência abaixo",
        "2. Abra o Gemini (gemini.google.com) na sua conta",
        "3. Cole o prompt + anexe a foto de uma peça genérica (still, fundo neutro)",
        "4. Gere a imagem",
        "5. Se ficou boa: salve o arquivo usando o 'nome_referencia' indicado",
        "6. Suba no Supabase Storage e marque a linha como 'aprovado' (ver README)",
        "",
        DSEP,
        "",
        "",
        HSEP,
        "# ABA: EVERGREEN",
        HSEP,
        "",
    ]

    for n, item in enumerate(evergreen, start=1):
        render_item(lines, n, f"{item['tipo_peca']} — {item['pose']}", item)

    lines += ["", HSEP, "# ABA: DATAS COMEMORATIVAS", HSEP, ""]

    for n, item in enumerate(datas, start=len(evergreen) + 1):
        titulo = f"{item['tipo_peca']} — {item['pose']} — {item['Data Comemorativa']}"
        render_item(lines, n, titulo, item)

    lines += ["", HSEP, "# ABA: EXPOSITORES", HSEP, ""]

    for n, item in enumerate(expositores, start=len(evergreen) + len(datas) + 1):
        render_item(lines, n, f"{item['tipo_peca']} — {item['pose']}", item)

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    total = len(evergreen) + len(datas) + len(expositores)
    print(f"OK: {total} prompts escritos em {out_path}")


if __name__ == "__main__":
    main()
