# -*- coding: utf-8 -*-
"""
Lê content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx e gera
supabase/seed.sql com um INSERT por linha das abas Evergreen e
Datas Comemorativas, na tabela reference_photos.

Todas as linhas entram com status = 'a_gerar' e thumbnail_url = NULL —
isso espelha o estado real de produção de conteúdo: a foto-mestre ainda
não existe. Assim que o time gerar e aprovar uma foto no Gemini, atualize
a linha correspondente (ver README.md, seção "Fluxo de conteúdo").

Uso:
    python3 scripts/seed_from_xlsx.py
"""
import openpyxl

XLSX_PATH = "content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx"
OUT_PATH = "supabase/seed.sql"


def esc(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def rows_from_evergreen(ws):
    out = []
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append({
            "categoria": "evergreen",
            "tipo_peca": row[idx["Tipo de Peça"]],
            "pose": row[idx["Pose/Enquadramento"]],
            "data_comemorativa": None,
            "nome_referencia": row[idx["Nome de Referência"]],
            "prompt_mestre": row[idx["Prompt Mestre"]],
            "perfis_sugeridos": row[idx["Perfis Sugeridos"]] or None,
            "aspecto": row[idx["Proporção"]],
            "ordem": row[idx["ID"]],
        })
    return out


def rows_from_datas(ws):
    out = []
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append({
            "categoria": "data_comemorativa",
            "tipo_peca": row[idx["Tipo de Peça"]],
            "pose": row[idx["Pose/Enquadramento"]],
            "data_comemorativa": row[idx["Data Comemorativa"]],
            "nome_referencia": row[idx["Nome de Referência"]],
            "prompt_mestre": row[idx["Prompt Mestre"]],
            "perfis_sugeridos": row[idx["Perfis Sugeridos"]] or None,
            "aspecto": row[idx["Proporção"]],
            "ordem": row[idx["ID"]],
        })
    return out


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    rows = rows_from_evergreen(wb["Evergreen"]) + rows_from_datas(wb["Datas Comemorativas"])

    lines = [
        "-- Gerado automaticamente por scripts/seed_from_xlsx.py",
        "-- Fonte: content/Photo_Studio_TPV_Biblioteca_de_Prompts.xlsx",
        "-- Não editar à mão — rode o script de novo após atualizar a planilha.",
        "",
    ]

    for r in rows:
        cols = [
            "categoria", "tipo_peca", "pose", "data_comemorativa",
            "nome_referencia", "prompt_mestre", "perfis_sugeridos",
            "aspecto", "ordem", "status",
        ]
        values = [
            esc(r["categoria"]), esc(r["tipo_peca"]), esc(r["pose"]),
            esc(r["data_comemorativa"]), esc(r["nome_referencia"]),
            esc(r["prompt_mestre"]), esc(r["perfis_sugeridos"]),
            esc(r["aspecto"]), str(r["ordem"]), "'a_gerar'",
        ]
        lines.append(
            f"insert into public.reference_photos ({', '.join(cols)}) "
            f"values ({', '.join(values)}) "
            f"on conflict (nome_referencia) do nothing;"
        )

    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    print(f"OK: {len(rows)} linhas escritas em {OUT_PATH}")


if __name__ == "__main__":
    main()
